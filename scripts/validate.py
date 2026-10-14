#!/usr/bin/env python
"""
Validiert die automatisch berechnete Spartenrechnung zellgenau gegen eine manuell
erstellte Referenz-Spartenrechnung (Soll-Ist-Vergleich je GuV-Zeile UND je Sparte).

Hintergrund: Die Summe einer Zeile kann zufaellig stimmen, waehrend die Verteilung
auf die Sparten falsch ist (genau dieser Fall ist bei den Frachtkosten aufgetreten).
Deshalb wird hier immer der komplette Vektor ueber alle Sparten geprueft.

Aufruf:
    python scripts/validate.py --rohdaten rohdaten/2026/Q1 \
        --referenz "referenz/01-03 2026-Spartenrechnung Perforator.xlsx" \
        --referenz-sheet "Spartenrechnung Jan-März 2026"
"""
import argparse
import os
import sys

import openpyxl

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.join(SCRIPT_DIR, "..")
CONFIG_DIR = os.path.join(REPO_ROOT, "config")
sys.path.insert(0, os.path.join(REPO_ROOT, "docs"))

import spartenrechnung_core as core  # noqa: E402
from spartenrechnung import find_kptm_file, find_kostentraeger_file  # noqa: E402

# Zuordnung: interner Zeilenschluessel -> Zeilennummer in der Referenz-Arbeitsmappe
REFERENZ_ZEILEN = {
    "erloese": 7,
    "bestand_fe": 8,
    "bestand_ufe": 9,
    "aktivierte_eigenleistung": 10,
    "betriebsleistung": 11,
    "material": 12,
    "fremdleistungen": 13,
    "personalaufwand": 14,
    "sbe": 16,
    "sba": 17,
    "afa": 18,
    "rohmarge_1": 20,
    "fek": 23,
    "fek_deckungsdifferenz": 24,
    "fek_nach_dd": 25,
    "db1": 26,
    "mgk": 28,
    "mgk_deckungsdifferenz": 29,
    "mgk_nach_dd": 30,
    "db2": 31,
    "vvgk": 33,
    "vvgk_deckungsdifferenz": 34,
    "vvgk_nach_dd": 35,
}

# Pseudo-Kostenobjekte (Sachkonten fuer Einmaleffekte/Abgrenzungen), die nicht aus
# der ERP-Produktgruppen-Dimension stammen und deshalb nicht verglichen werden.
PSEUDO_KOSTENOBJEKTE = {"712000", "721000", "910000", "920000", "999999"}

TOLERANZ_EXAKT = 0.005

# Von der Produktgruppen-Reorganisation 2026 betroffen: nur die Pressbohrtechnik-
# Gruppen (ohne das Fracht-Kostenobjekt PB0100, das bleibt). Die Referenz ist in der
# alten Struktur gerechnet - in der neuen ist deshalb nur ihre SUMME vergleichbar.
def ist_reorganisiert(code):
    return code.startswith("PB") and code != "PB0100"


FAMILIE = "PB (reorg.)"


def zusammenfassen(werte_je_pg):
    aus = {}
    for pg, w in werte_je_pg.items():
        k = FAMILIE if ist_reorganisiert(pg) else pg
        aus[k] = aus.get(k, 0.0) + float(w or 0.0)
    return aus


def lade_referenz(pfad, sheet):
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb[sheet]
    spalten = {}
    for c in range(4, 40):
        code = ws.cell(row=4, column=c).value
        if code is not None:
            spalten[str(code).strip()] = c
    werte = {}
    for key, zeile in REFERENZ_ZEILEN.items():
        werte[key] = {
            code: (ws.cell(row=zeile, column=c).value or 0.0)
            for code, c in spalten.items()
            if code not in PSEUDO_KOSTENOBJEKTE
        }
    return werte


def main():
    ap = argparse.ArgumentParser(description="Vergleicht die automatische Spartenrechnung zellgenau mit einer manuellen Referenz.")
    ap.add_argument("--rohdaten", required=True)
    ap.add_argument("--referenz", required=True)
    ap.add_argument("--referenz-sheet", required=True)
    ap.add_argument("--zeitraum", help="Fuer den Jahresbezug des Werkstattbestands, z.B. 'Jan-Maerz 2026'")
    ap.add_argument("--pwbs-ende", help="Optional: Werkstattbestand zum Periodenende (fuer die UFE-Zeile)")
    ap.add_argument("--pfak", nargs="*", help="Optional: PFAK-Dateien zur Sparten-Zuordnung")
    ap.add_argument("--kostentraeger", help="Optional: Kostentraeger-Auswertung (Deckungsdifferenzen). "
                    "Ohne Angabe wird im Rohdaten-Ordner gesucht.")
    ap.add_argument("--alte-struktur", action="store_true",
                    help="Ohne Umschluesselung rechnen (wie gebucht). Dann ist auch PB je "
                         "Sparte vergleichbar - die Referenz ist in der alten Struktur gerechnet.")
    args = ap.parse_args()

    mapping, pg_config = core.load_config(CONFIG_DIR)
    df = core.load_kptm(find_kptm_file(args.rohdaten))
    satz = core.ermittle_stundensatz(df, mapping)
    if satz:
        mapping["standard_stundensatz"] = satz

    # Gleiche Reihenfolge wie generate(): erst umschluesseln, dann rechnen.
    params = core.lade_betriebsparameter(CONFIG_DIR)
    jahr = core.geschaeftsjahr(df, args.zeitraum or args.referenz_sheet)
    karte = {} if args.alte_struktur else (params.get("umschluesselung") or {})
    aufgeloest = pg_config.get("aufgeloest", {}) if karte else {}
    if karte:
        df, bef = core.umschluesseln(df, karte, aufgeloest)
        print(f"Neue Struktur: {bef['zeilen']} Buchungszeilen umgeschluesselt. PB-Gruppen werden "
              f"als Summe '{FAMILIE}' verglichen.\n")
    result, produktgruppen, _, _, _ = core.build_spartenrechnung(df, mapping, pg_config)

    kt_pfad = args.kostentraeger or find_kostentraeger_file(args.rohdaten)
    if kt_pfad:
        result, dd = core.deckungsdifferenzen_einsetzen(
            result, kt_pfad, mapping, df,
            anteile=core.aufteilungsanteile(df, mapping) if karte else None)
        for h in dd.get("hinweise", []):
            print(f"HINWEIS: {h}", file=sys.stderr)

    if args.pwbs_ende:
        ab = core.lade_anfangsbestand(CONFIG_DIR, jahr)
        if args.alte_struktur:
            print("HINWEIS: --alte-struktur mit einem Anfangsbestand in neuer Struktur - die "
                  "UFE-Zeile der PB-Gruppen ist dann nicht vergleichbar.\n", file=sys.stderr)
        ufe, warnungen, _stat = core.bestandsveraenderung_ufe(
            args.pwbs_ende, args.pfak, anfangsbestand=ab,
            jahr=jahr, kptm_df=df, karte=karte, aufgeloest=aufgeloest)
        for w in warnungen:
            print(f"HINWEIS: {w}\n", file=sys.stderr)
        if ufe:
            for pg, wert in ufe.items():
                if pg in result.columns:
                    result.loc["bestand_ufe", pg] = wert
            result = core.neu_berechnen(result)

    referenz = lade_referenz(args.referenz, args.referenz_sheet)
    if karte:
        referenz = {k: zusammenfassen(v) for k, v in referenz.items()}
        ist_je_zeile = {k: zusammenfassen({pg: result.loc[k, pg] for pg in produktgruppen})
                        for k in REFERENZ_ZEILEN if k in result.index}
        produktgruppen = list(dict.fromkeys(
            FAMILIE if ist_reorganisiert(pg) else pg for pg in produktgruppen))
    else:
        ist_je_zeile = {k: {pg: result.loc[k, pg] for pg in produktgruppen}
                        for k in REFERENZ_ZEILEN if k in result.index}

    labels = mapping["zeilen_labels"]
    print(f"{'Zeile':42s} {'exakt':>6s} {'nah':>5s} {'ab':>4s} {'max Δ':>12s}  {'Summe Δ':>13s}")
    print("-" * 92)

    gesamt_exakt = gesamt_nah = gesamt_ab = 0
    problemzeilen = []

    for key in REFERENZ_ZEILEN:
        if key not in result.index:
            continue
        soll_map = referenz[key]
        exakt = nah = ab = 0
        maxdiff = 0.0
        soll_summe = ist_summe = 0.0
        abweichungen = []
        for pg in produktgruppen:
            soll = float(soll_map.get(pg, 0.0) or 0.0)
            ist = float(ist_je_zeile[key].get(pg, 0.0))
            soll_summe += soll
            ist_summe += ist
            if abs(soll) < TOLERANZ_EXAKT and abs(ist) < TOLERANZ_EXAKT:
                continue
            diff = ist - soll
            maxdiff = max(maxdiff, abs(diff))
            if abs(diff) < TOLERANZ_EXAKT:
                exakt += 1
            elif abs(diff) < max(1.0, abs(soll) * 0.01):
                nah += 1
            else:
                ab += 1
                abweichungen.append((pg, soll, ist))
        gesamt_exakt += exakt
        gesamt_nah += nah
        gesamt_ab += ab
        marker = "" if ab == 0 else "  <-- pruefen"
        print(f"{labels.get(key, key)[:42]:42s} {exakt:6d} {nah:5d} {ab:4d} {maxdiff:12,.2f}  {ist_summe - soll_summe:+13,.2f}{marker}")
        if abweichungen:
            problemzeilen.append((key, abweichungen))

    print("-" * 92)
    print(f"{'GESAMT':42s} {gesamt_exakt:6d} {gesamt_nah:5d} {gesamt_ab:4d}")
    print("\nexakt = centgenau | nah = <1 % bzw. <1 EUR Abweichung (Buchungsstand/Rundung) | ab = echte Abweichung")

    if problemzeilen:
        print("\nDetails zu den abweichenden Zellen:")
        for key, abweichungen in problemzeilen:
            print(f"\n  {labels.get(key, key)}:")
            for pg, soll, ist in abweichungen[:20]:
                print(f"    {pg}: Referenz {soll:14,.2f} | berechnet {ist:14,.2f} | Δ {ist - soll:+12,.2f}")


if __name__ == "__main__":
    main()
